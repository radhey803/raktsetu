"""
build_raktsetu_docx.py
Project report generator for 'Raktsetu' in the CMPICA (CHARUSAT) report format.
Institution: CMPICA, FCA, CHARUSAT
Authors: Radhey Gajera (25MCA057), Shruti Gajipara (25MCA058)
Guide: Ms. Konica Soni (Assistant Professor)
Submission: November 2026

Layout follows the department template: bordered pages, project name in the header,
"Page N" footer, section divider pages, ❖/✓ heading markers and Times New Roman text.
Every topic starts on a new page.

The Table of Contents page numbers are computed, not typed: the report is built once,
rendered to PDF with LibreOffice, each heading is located in the PDF text, and the
report is rebuilt with those numbers. Usage: uv run python scripts/report/build_raktsetu_docx.py
"""

import os
import shutil
import subprocess
import tempfile

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
ASSETS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets")
DIAGRAMS_DIR = os.path.join(REPO_ROOT, "docs", "diagrams")
SCREENSHOTS_DIR = os.path.join(REPO_ROOT, "docs", "screenshots")
OUTPUT_FILE = os.path.join(REPO_ROOT, "reports", "Raktsetu_Project_Report.docx")

PROJECT = "Raktsetu"
FONT = "Times New Roman"
RED = RGBColor(0xFF, 0x00, 0x00)
GREEN = RGBColor(0x00, 0xA8, 0x4F)
BLACK = RGBColor(0, 0, 0)
STUDENTS = ["Radhey Gajera (25MCA057)", "Shruti Gajipara (25MCA058)"]
GUIDE = "Ms. Konica Soni"
SUBMISSION = "November 2026"


# ============================================================================
# Low-level formatting helpers
# ============================================================================
def _font(run, size=12, bold=False, italic=False, underline=False, color=BLACK, name=FONT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.underline = underline
    run.font.color.rgb = color
    return run


def para(
    doc,
    text="",
    size=12,
    bold=False,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    color=BLACK,
    before=0,
    after=6,
    indent=0.0,
    underline=False,
    italic=False,
):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = 1.15
    if indent:
        pf.left_indent = Inches(indent)
    if text:
        _font(p.add_run(text), size, bold, italic, underline, color)
    return p


def new_page(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def divider(doc, title):
    """Full page with only the chapter title, centred (department template)."""
    new_page(doc)
    para(
        doc, title, size=30, bold=True, underline=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=290
    )


def heading(doc, text, page_break=True):
    """Topic heading, centred at the top of a new page."""
    if page_break:
        new_page(doc)
    return para(
        doc,
        text,
        size=16,
        bold=True,
        underline=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=4,
        after=14,
    )


def subheading(doc, text, indent=0.25, size=13):
    p = para(doc, "", align=WD_ALIGN_PARAGRAPH.LEFT, before=8, after=4, indent=indent)
    _font(p.add_run("❖ "), size, bold=True)
    _font(p.add_run(text), size, bold=True)
    return p


def check(doc, label, text=None, indent=0.6):
    """✓ Label: text  (bold label, normal text)."""
    p = para(doc, "", align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=3, indent=indent)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    _font(p.add_run("✓ "), 12, bold=True)
    _font(p.add_run(label + (": " if text else "")), 12, bold=bool(text) or len(label) < 45)
    if text:
        _font(p.add_run(text), 12)
    return p


def bullet(doc, text, label=None, indent=0.95):
    p = para(doc, "", after=2, indent=indent)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    _font(p.add_run("•  "), 12)
    if label:
        _font(p.add_run(label + ": "), 12, bold=True)
    _font(p.add_run(text), 12)
    return p


def numbered(doc, n, title, text, indent=0.6):
    p = para(doc, "", after=1, indent=indent)
    _font(p.add_run(f"{n}. {title}"), 12, bold=True)
    para(doc, text, indent=indent + 0.45, after=4)


def tool(doc, name, purpose, features, benefits):
    check(doc, name)
    bullet(doc, purpose, "Purpose")
    bullet(doc, features, "Features")
    bullet(doc, benefits, "Benefits")


def _set_widths(t, widths):
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = Inches(w)
        for cell in t.columns[i].cells:
            cell.width = Inches(w)


def _shade(cell, hex_color):
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'))


FIGS, TABS = [], []  # captions collected while building, numbered in order


def fig_caption(doc, title):
    FIGS.append(title)
    para(
        doc,
        f"Figure {len(FIGS)}: {title}",
        size=11,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=10,
    )


def table_caption(doc, title):
    TABS.append(title)
    para(
        doc,
        f"Table {len(TABS)}: {title}",
        size=11,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=6,
        after=4,
    )


def table(
    doc, headers, rows, widths, header_fill="D9D9D9", font_size=11, center_cols=(), caption=None
):
    if caption:
        table_caption(doc, caption)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        _shade(c, header_fill)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(c.paragraphs[0].add_run(h), font_size, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in center_cols else WD_ALIGN_PARAGRAPH.LEFT
            if value == "PASS":
                _font(p.add_run("✔"), font_size + 3, bold=True, color=GREEN)
            else:
                _font(
                    p.add_run(str(value)), font_size, bold=(i == 1 and headers[0].startswith("Sr"))
                )
    _set_widths(t, widths)
    para(doc, "", after=4)
    return t


def image(doc, path, width=6.2, max_height=7.2, caption=None):
    if not os.path.exists(path):
        para(doc, f"[missing image: {os.path.basename(path)}]", color=RED)
        return
    from PIL import Image

    with Image.open(path) as im:
        w, h = im.size
    width = min(width, max_height * w / h)  # keep tall images on one page
    p = para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=2)
    p.add_run().add_picture(path, width=Inches(width))
    if caption:
        fig_caption(doc, caption)


def images_side_by_side(doc, paths, width_each=3.0, caption=None):
    t = doc.add_table(rows=1, cols=len(paths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, path in zip(t.rows[0].cells, paths, strict=True):
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run().add_picture(path, width=Inches(width_each))
    if caption:
        fig_caption(doc, caption)


def screen(doc, title, description, shot, width=6.2):
    heading(doc, title)
    check(doc, description)
    image(doc, os.path.join(SCREENSHOTS_DIR, shot), width=width, max_height=6.9, caption=title)


def diagram(doc, title, text, file, max_height=7.0):
    heading(doc, title)
    check(doc, text)
    image(doc, os.path.join(DIAGRAMS_DIR, file), width=6.3, max_height=max_height, caption=title)


# ============================================================================
# Page setup: border, header, footer
# ============================================================================
def setup_document(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(12)

    sec = doc.sections[0]
    sec.page_height, sec.page_width = Inches(11), Inches(8.5)
    sec.top_margin = sec.bottom_margin = Inches(1)
    sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.55)

    # Thick-thin double border around every page, as in the department template
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "thickThinSmallGap")
        b.set(qn("w:sz"), "24")
        b.set(qn("w:space"), "24")
        b.set(qn("w:color"), "000000")
        borders.append(b)
    sec._sectPr.find(qn("w:pgMar")).addnext(borders)

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _font(hp.add_run(PROJECT), 12, color=RED)

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(fp.add_run("Page "), 11, name="Calibri")
    field = parse_xml(
        f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"><w:r><w:rPr>'
        f'<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:sz w:val="22"/></w:rPr>'
        f"<w:t>1</w:t></w:r></w:fldSimple>"
    )
    fp._p.append(field)


# ============================================================================
# Front matter
# ============================================================================
def cover(doc):
    c = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, PROJECT, size=28, color=RED, align=c, before=30, after=2)
    para(
        doc,
        "eRaktKosh Blood Lakehouse: National Blood Availability & Shortage Intelligence",
        size=13,
        italic=True,
        align=c,
        after=8,
    )
    para(doc, "By", size=16, align=c, after=4)
    for s in STUDENTS:
        para(doc, s, size=18, color=RED, align=c, after=2)
    para(doc, "Submitted to", size=16, align=c, before=36, after=2)
    para(doc, GUIDE, size=18, color=RED, align=c, after=0)
    para(doc, "Assistant Professor, CMPICA", size=12, align=c, after=6)
    logo = os.path.join(ASSETS_DIR, "cmpica_logo.jpeg")
    if os.path.exists(logo):
        para(doc, "", align=c, after=4).add_run().add_picture(logo, width=Inches(1.25))
    para(
        doc,
        "Smt. Chandaben Mohanbhai Patel Institute of Computer Applications",
        size=14,
        align=c,
        after=0,
    )
    para(doc, "CHARUSAT", size=14, align=c, after=0)
    para(doc, "Changa", size=14, align=c, after=30)
    para(doc, SUBMISSION, size=16, color=RED, align=c, after=4)
    logo = os.path.join(ASSETS_DIR, "charusat_logo.jpeg")
    if os.path.exists(logo):
        para(doc, "", align=c, after=2).add_run().add_picture(logo, width=Inches(3.4))
    para(doc, "Accredited with Grade A+ by NAAC,", size=11, align=c, after=0)
    para(doc, "CHAROTAR UNIVERSITY OF SCIENCE AND TECHNOLOGY", size=16, align=c, after=0)
    para(doc, "Changa", size=16, align=c, after=0)


def declaration(doc):
    new_page(doc)
    para(
        doc,
        "Declaration",
        size=16,
        bold=True,
        underline=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=30,
        after=18,
    )
    para(
        doc,
        f"We, {STUDENTS[0]} and {STUDENTS[1]}, students of Master of Computer Applications (MCA) at "
        "Smt. Chandaben Mohanbhai Patel Institute of Computer Applications (CMPICA), CHARUSAT, declare that "
        "the project report titled \u201cRaktsetu: eRaktKosh Blood Lakehouse\u201d is our own work. It was "
        f"carried out under the guidance of {GUIDE}.",
        after=12,
    )
    para(
        doc,
        "This work has not been submitted to any other university or institute for any degree or diploma. "
        "All data sources, tools and references used in this project are acknowledged in this report.",
        after=40,
    )
    para(doc, "Place: Changa", align=WD_ALIGN_PARAGRAPH.LEFT, after=4)
    para(doc, "Date: ____________", align=WD_ALIGN_PARAGRAPH.LEFT, after=40)
    for s_ in STUDENTS:
        para(doc, "____________________", align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
        para(doc, s_, align=WD_ALIGN_PARAGRAPH.RIGHT, after=18)


def abstract(doc):
    new_page(doc)
    para(
        doc,
        "Abstract",
        size=16,
        bold=True,
        underline=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=30,
        after=18,
    )
    for text in [
        "When a patient needs blood urgently, families often phone one blood bank after another. India's eRaktKosh "
        "portal already publishes blood stock for thousands of blood centres, but it is searched one district at "
        "a time, gives no help with blood-group compatibility and does not show how old each record is.",
        "Raktsetu solves this by building a data lakehouse on top of eRaktKosh. An async scraper collects stock "
        "for every district and saves it as immutable Parquet files. DuckDB and dbt clean the data in Bronze, "
        "Silver and Gold layers, keep a daily history and run 37 data-quality tests. Only data that passes these "
        "checks is published to a read-only database, which a FastAPI REST API serves to a responsive website.",
        "The website offers a Smart Emergency Mode that lists compatible donor groups (including plasma and the "
        "rare Bombay group) and ranks nearby centres by freshness, trust, units and distance. It also provides a "
        "blood finder with a live map, 7-day trends, a Trust Score, rare-blood and district-shortage radars, CSV "
        "export, a donor eligibility checker and a data-health dashboard. It works on any phone and offline.",
        "The system covers 4,716 blood centres in 35 States/UTs and 706 districts. It is supported by 98 automated "
        "code tests, 37 data tests and 146 end-to-end browser checks, all passing.",
    ]:
        para(doc, text, after=10)
    p = para(doc, "", after=0, before=6)
    _font(p.add_run("Keywords: "), 12, bold=True)
    _font(
        p.add_run(
            "blood availability, eRaktKosh, data lakehouse, DuckDB, dbt, FastAPI, Progressive Web App, "
            "emergency healthcare"
        ),
        12,
        italic=True,
    )


def list_page(doc, title, kind, items, pages):
    heading(doc, title)
    rows = [(f"{kind} {n}", t, pages.get(f"{kind} {n}:", "00")) for n, t in enumerate(items, 1)]
    table(
        doc,
        [f"{kind} No.", "Title", "Page"],
        rows,
        (1.0, 4.4, 0.8),
        font_size=11,
        center_cols=(0, 2),
    )


def project_timeline(doc):
    heading(doc, "Project Timeline (Gantt Chart)")
    check(
        doc,
        "The project was planned in eight phases over sixteen weeks. Some phases overlap, for example the "
        "website was built while the API was being finished.",
    )
    image(
        doc,
        os.path.join(DIAGRAMS_DIR, "gantt_chart.png"),
        width=6.3,
        caption="Project timeline (Gantt chart)",
    )


def comparison(doc):
    heading(doc, "Existing vs Proposed System")
    rows = [
        ("Search scope", "One district at a time", "Whole country, any filter"),
        ("Blood compatibility help", "No", "Yes (red cells, plasma, Bombay)"),
        ("Nearest blood centre", "No", "Yes, using device location"),
        ("Shows how fresh the data is", "Partly (date only)", "Yes, on every result"),
        ("Trust / reliability score", "No", "Yes (0–100)"),
        ("Stock history and trends", "No", "Yes (daily snapshots)"),
        ("District shortage analysis", "No", "Yes, with severity score"),
        ("Rare (Bombay) blood search", "No", "Yes, nationwide"),
        ("Report export", "No", "Yes, CSV"),
        ("Donor eligibility check", "No", "Yes, private, in browser"),
        ("Works offline", "No", "Yes (PWA)"),
        ("Open API for other apps", "No", "Yes (REST API)"),
    ]
    table(
        doc,
        ["Feature", "Existing (eRaktKosh)", "Proposed (Raktsetu)"],
        rows,
        (2.4, 1.9, 2.0),
        caption="Comparison of existing and proposed system",
    )


def appendix(doc):
    divider(doc, "APPENDIX")
    heading(doc, "Appendix A: Glossary")
    table(
        doc,
        ["Term", "Meaning"],
        [
            ("API", "Application Programming Interface: a way for programs to request data"),
            ("Bronze / Silver / Gold", "Raw, cleaned and analysis-ready layers of the lakehouse"),
            ("Bombay phenotype (Oh)", "A very rare blood group that can only receive Bombay blood"),
            ("CSV", "Comma-Separated Values: a simple table file that opens in Excel"),
            ("dbt", "Data Build Tool: turns SQL files into tested tables"),
            ("DFD", "Data Flow Diagram"),
            ("DuckDB", "A fast analytical database stored in a single file"),
            ("ELT", "Extract, Load, Transform: the order in which the pipeline handles data"),
            ("FFP", "Fresh Frozen Plasma"),
            ("Haversine formula", "Formula for the distance between two points on Earth"),
            ("Lakehouse", "Data design that keeps raw files and clean tables together"),
            ("NBTC", "National Blood Transfusion Council, Govt. of India"),
            ("Parquet", "A compressed, column-based file format for data"),
            ("PRBC", "Packed Red Blood Cells"),
            ("PWA", "Progressive Web App: a website that can be installed and used offline"),
            ("Rh-negative", "Blood without the RhD antigen (e.g. O-, A-)"),
            ("Trust Score", "Our 0–100 score of how reliable a centre's data is"),
            ("UML", "Unified Modeling Language, used for the diagrams"),
        ],
        (2.0, 4.3),
        caption="Glossary of terms",
    )

    heading(doc, "Appendix B: User Manual")
    subheading(doc, "For Public Users", indent=0.0)
    for n, (t, d) in enumerate(
        [
            (
                "Emergency search",
                "Open the website, choose the patient's blood group, component, state and district, "
                "then press the emergency search button. Call the top centre to confirm.",
            ),
            (
                "Near me",
                "Press 'Use my location' or 'Near me' and allow location access to sort by distance.",
            ),
            (
                "Find blood",
                "In Find Blood, choose filters and press Search. Switch to the map view if needed.",
            ),
            ("Trend and trust", "On any card, press '7-day trend' or the Trust badge for details."),
            ("Reports", "Press 'Export CSV' in Find Blood or Shortages to download a report."),
            (
                "Donor check",
                "In Donor Help, fill in age, weight and the questions, then press Evaluate.",
            ),
        ],
        1,
    ):
        numbered(doc, n, t, d)
    subheading(doc, "For the Operator", indent=0.0)
    for n, (t, d) in enumerate(
        [
            ("Install", "Install Python 3.12 and uv, then run: make install"),
            ("Refresh data", "Run: make pipeline (scrape, load, transform, test and publish)"),
            ("Start the website", "Run: make run, then open http://localhost:8000"),
            ("Check data health", "Open the Data Health section or run: make status"),
            ("Rebuild this report", "Run: make report"),
        ],
        1,
    ):
        numbered(doc, n, t, d)


def acknowledgement(doc):
    new_page(doc)
    para(
        doc,
        "Acknowledgement",
        size=16,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=30,
        after=18,
    )
    for text in [
        "Knowledge in itself is a continuous process. At this moment of our substantial enhancement, we rarely "
        "find words to express our gratitude towards those who were constantly involved with us.",
        "The completion of any inter-disciplinary project depends upon coordination, cooperation and the combined "
        "efforts of several resources of knowledge, creativity, skill, energy and time. With the work now "
        "accomplished, we take this opportunity to acknowledge, in these lines, everyone who contributed to it.",
    ]:
        para(doc, text, after=12)
    p = para(doc, "We would like to thank our project guide ", after=12)
    _font(p.add_run(GUIDE), 12, underline=True)
    _font(p.add_run(", I/C Principal "), 12)
    _font(p.add_run("Dr. Dharmendra Patel"), 12, bold=True)
    _font(p.add_run(" and I/C Dean "), 12)
    _font(p.add_run("Dr. Sanskruti Patel"), 12, bold=True)
    _font(
        p.add_run(
            " who advised and gave us moral support throughout the duration of our project. Without "
            "their constant encouragement we could not have achieved what we have."
        ),
        12,
    )
    para(
        doc,
        "We also thank the Ministry of Health and Family Welfare (MoHFW) for publishing blood stock data "
        "openly through the eRaktKosh portal, and the open-source communities behind DuckDB, dbt, "
        "FastAPI, Leaflet and OpenStreetMap, whose work made this project possible.",
        after=12,
    )
    para(
        doc,
        "It is our good fortune that we had the support and well wishes of many. We are thankful to all, "
        "including those whose names may have been missed here; their contributions have not gone unnoticed.",
        after=30,
    )
    para(doc, "With Sincere Regards,", align=WD_ALIGN_PARAGRAPH.RIGHT, after=12)
    for s in STUDENTS:
        para(doc, s, color=RED, align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)


# (sr_no, sub_no, subject, anchor text searched in the PDF, bold row)
TOC = [
    ("", "", "List of Figures", "List of Figures", True),
    ("", "", "List of Tables", "List of Tables", True),
    ("1", "", "Project Profile", "PROJECT PROFILE", True),
    ("", "", "Project Timeline (Gantt Chart)", "Project Timeline (Gantt Chart)", False),
    ("2", "", "Introduction to Tools", "INTRODUCTION TO TOOLS", True),
    ("3", "", "System Study", "SYSTEM STUDY", True),
    ("", "1.", "Existing System", "Existing System", False),
    ("", "2.", "Proposed System", "Proposed System", False),
    ("", "3.", "Existing vs Proposed System", "Existing vs Proposed System", False),
    ("", "4.", "Scope of the Proposed System", "Scope of the Proposed System", False),
    ("", "5.", "Aim and Objective of the Proposed System", "Aim and Objectives", False),
    ("", "6.", "Feasibility Study", "Feasibility Study", False),
    ("", "", "Operational Feasibility", "Operational Feasibility", False),
    ("", "", "Technical Feasibility", "Technical Feasibility", False),
    ("", "", "Economical Feasibility", "Economical Feasibility", False),
    ("4", "", "System Analysis", "SYSTEM ANALYSIS", True),
    (
        "",
        "",
        "Requirements Specification (along with System Modules)",
        "Requirement Specification (along with System Modules)",
        False,
    ),
    ("", "", "Hardware and Software Requirements", "Hardware and Software Requirements", False),
    ("", "", "Data Flow Diagram (Level 0)", "Data Flow Diagram (Level 0)", False),
    ("", "", "Data Flow Diagram (Level 1)", "Data Flow Diagram (Level 1)", False),
    ("", "", "Use Case Diagram", "Use Case Diagram: Public Users", False),
    ("", "", "Class Diagram", "Class Diagram", False),
    ("", "", "Activity Diagram", "Activity Diagram: Public User", False),
    ("", "", "Sequence Diagram", "Sequence Diagram", False),
    ("", "", "System Architecture", "System Architecture", False),
    ("5", "", "System Design", "SYSTEM DESIGN", True),
    ("", "", "ER Diagram", "ER Diagram", False),
    ("", "", "Data Dictionary", "Data Dictionary", False),
    ("", "", "Screen Layouts", "Screen Layouts", False),
    ("6", "", "System Testing", "SYSTEM TESTING", True),
    ("", "", "Testing Strategies", "Testing Strategies", False),
    ("", "", "Test Results Summary", "Test Results Summary", False),
    ("", "", "Test Cases", "Test Cases", False),
    ("7", "", "Conclusion and Limitations", "CONCLUSION", True),
    ("8", "", "Future Enhancement", "FUTURE ENHANCEMENT", True),
    ("9", "", "Bibliography/References", "REFERENCES", True),
    ("10", "", "Appendix", "APPENDIX", True),
    ("", "A.", "Glossary", "Appendix A: Glossary", False),
    ("", "B.", "User Manual", "Appendix B: User Manual", False),
    ("", "", "Reporting Report (to be collected from respective internal guide)", None, True),
]


def table_of_contents(doc, pages):
    new_page(doc)
    para(doc, "Table of Contents", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["Sr. No", "", "Subject", "Page Number"]):
        _font(t.rows[0].cells[i].paragraphs[0].add_run(h), 12, bold=True)
    for sr, sub, subject, anchor, bold in TOC:
        cells = t.add_row().cells
        _font(cells[0].paragraphs[0].add_run(sr), 12, bold=True)
        cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _font(cells[1].paragraphs[0].add_run(sub), 12, bold=True)
        _font(cells[2].paragraphs[0].add_run(subject), 12, bold=bold)
        page = str(pages.get(subject, "00")) if anchor else ""
        _font(cells[3].paragraphs[0].add_run(page), 12, bold=True)
    _set_widths(t, (0.75, 0.55, 3.6, 1.4))


# ============================================================================
# 1. Project profile
# ============================================================================
def project_profile(doc):
    divider(doc, "PROJECT PROFILE")
    heading(doc, "Project Profile")
    rows = [
        ("Project Name", "Raktsetu (eRaktKosh Blood Lakehouse)"),
        (
            "Type of Application",
            "Web Application (Progressive Web App) with a Data Engineering Pipeline",
        ),
        (
            "Project Description",
            "Raktsetu turns the blood-stock data published by the Government of India's eRaktKosh portal into a "
            "fast, honest emergency finder. A medallion data lakehouse collects nationwide stock from every "
            "district, cleans and tests it, keeps a daily history, and publishes it to a responsive website. "
            "Patients' families can find compatible blood near them in seconds, see how fresh each record is, "
            "call or navigate to the blood centre, and officials can spot district-level shortages and export "
            "them as reports.",
        ),
        (
            "Data Source",
            "eRaktKosh (Ministry of Health and Family Welfare, Govt. of India), public portal",
        ),
        ("Coverage", "4,716 blood centres across 35 States/UTs and 706 districts"),
        ("Team Size", "2"),
        ("Frontend", "HTML5, CSS3, Vanilla JavaScript, Leaflet.js (maps), Chart.js (charts)"),
        ("Backend", "Python 3.12, FastAPI, Uvicorn"),
        ("Data Pipeline", "Async scraper (aiohttp), Apache Parquet, dbt Core (dbt-duckdb)"),
        ("Database", "DuckDB: analytical warehouse plus a read-only serving replica"),
        ("Tools used", "VS Code, uv, Git & GitHub, Docker, pytest, Playwright, Ruff"),
        ("Internal Guide", f"{GUIDE}, Assistant Professor, CMPICA"),
    ]
    for label, text in rows:
        p = para(doc, "", after=8, indent=0.3)
        _font(p.add_run(label + ": "), 12, bold=True)
        _font(p.add_run(text), 12)


# ============================================================================
# 2. Introduction to tools
# ============================================================================
def introduction_to_tools(doc):
    divider(doc, "INTRODUCTION TO TOOLS")
    heading(doc, "Introduction to Tools & Technology")
    subheading(doc, "Front End:")
    tool(
        doc,
        "HTML5 & CSS3",
        "Structure and visual design of every screen of the website.",
        "Semantic sections, CSS Grid/Flexbox layouts, media queries, design tokens.",
        "Fully responsive from 320 px phones to 1920 px desktops without any UI framework.",
    )
    tool(
        doc,
        "JavaScript (ES2020)",
        "All client-side behaviour: search, filters, modals, map and charts.",
        "Fetch API, async/await, URL state sync, Geolocation API, client-side eligibility rules.",
        "Zero build step and a small download, so the site loads fast on slow mobile networks.",
    )
    tool(
        doc,
        "Leaflet.js",
        "Interactive map of blood centres.",
        "Custom markers, pop-ups, fit-to-results and a 'you are here' marker.",
        "Lightweight open-source mapping that works well on mobile devices.",
    )
    tool(
        doc,
        "Chart.js",
        "National distribution chart and per-bank historical trend chart.",
        "Bar and line charts with tooltips; gaps are shown for days without data.",
        "Clear visual insight with very little code.",
    )
    tool(
        doc,
        "Service Worker & Web App Manifest (PWA)",
        "Installable app with offline support.",
        "Network-first caching of the app shell, libraries and last API responses.",
        "The last known inventory stays visible during network outages.",
    )

    heading(doc, "Back End & Data Engineering")
    subheading(doc, "Back End:")
    tool(
        doc,
        "Python 3.12",
        "Primary language for the API, scraper, pipeline CLI and tests.",
        "Type hints, asyncio, rich standard library.",
        "Readable, maintainable code with a large data ecosystem.",
    )
    tool(
        doc,
        "FastAPI & Uvicorn",
        "REST API serving the website and third-party clients.",
        "Automatic OpenAPI docs (/docs), request validation with Pydantic, CSV export endpoints.",
        "High performance with self-documenting, type-checked endpoints.",
    )
    subheading(doc, "Data Engineering:")
    tool(
        doc,
        "aiohttp (Async Scraper)",
        "Collects blood stock for every district from eRaktKosh.",
        "10 concurrent requests, retries with exponential back-off, circuit breaker at 20% errors.",
        "A full national crawl in minutes while staying polite to the government server.",
    )
    tool(
        doc,
        "Apache Parquet",
        "Immutable landing zone for every raw scrape.",
        "Columnar, compressed files partitioned by date and hour with a SHA-256 manifest.",
        "Raw data is never lost and any run can be reloaded or audited.",
    )
    tool(
        doc,
        "DuckDB",
        "Embedded analytical database (warehouse and serving copy).",
        "Columnar SQL engine, reads Parquet directly, single-file database.",
        "Analytics in milliseconds with no database server to install or manage.",
    )
    tool(
        doc,
        "dbt Core (dbt-duckdb)",
        "Transforms raw data into clean Silver and Gold tables.",
        "Version-controlled SQL models, incremental snapshots, 37 automated data tests, lineage docs.",
        "Reliable, testable and documented transformations.",
    )

    heading(doc, "Key Services & Architecture")
    subheading(doc, "Key Services:")
    tool(
        doc,
        "eRaktKosh Public Portal",
        "Single source of truth for blood stock in India.",
        "District-wise availability per blood group and component for registered blood centres.",
        "Official, nationwide data; every page and export credits MoHFW as the source.",
    )
    tool(
        doc,
        "OpenStreetMap Nominatim",
        "Geocoding of all 706 districts to map coordinates.",
        "One-time lookup at 1 request/second, cached in data/district_coords.json.",
        "Accurate 'Near me' distances without a paid geocoding service.",
    )
    tool(
        doc,
        "CARTO Positron Basemap",
        "Muted background map tiles.",
        "Light, low-contrast tiles that keep blood centre markers easy to read.",
        "Clean, professional map appearance.",
    )
    subheading(doc, "Architecture:")
    tool(
        doc,
        "Medallion Lakehouse (Bronze, Silver, Gold)",
        "Layered data design.",
        "Bronze keeps raw rows, Silver cleans and parses, Gold holds analysis-ready marts.",
        "Each layer is testable and every number can be traced back to its raw source.",
    )
    tool(
        doc,
        "Atomic Blue/Green Publishing",
        "Safe release of new data to the website.",
        "An anomaly gate checks the data, then the warehouse is copied to serving.duckdb atomically.",
        "Users never see half-updated or broken data.",
    )

    heading(doc, "Development & Testing Tools")
    tool(
        doc,
        "uv & Ruff",
        "Dependency management, linting and formatting.",
        "Locked, reproducible environments; fast lint and format checks.",
        "Consistent code quality across the team.",
    )
    tool(
        doc,
        "pytest",
        "Unit and integration testing of Python code.",
        "98 automated tests covering scraper, scoring, geocoding, CSV export and every API endpoint.",
        "Regressions are caught before they reach users.",
    )
    tool(
        doc,
        "Playwright",
        "Automated browser (end-to-end) testing.",
        "Drives the real website on desktop and mobile screen sizes; 146 user-flow checks.",
        "Proves that features work from the user's side, not just in code.",
    )
    tool(
        doc,
        "Docker & Git",
        "Packaging and version control.",
        "Container image with health check; Git history on GitHub.",
        "The same build runs on any machine; all changes are tracked.",
    )


# ============================================================================
# 3. System study
# ============================================================================
def system_study(doc):
    divider(doc, "SYSTEM STUDY")
    heading(doc, "Existing System")
    para(doc, "Existing System: eRaktKosh Portal (MoHFW, Govt. of India)", bold=True, after=6)
    for label, text in [
        (
            "District search",
            "Users select a state, district, blood group and component to view stock.",
        ),
        ("Facility details", "Shows the blood centre name, contact details and last update time."),
        ("Government data", "Stock is entered by registered blood centres across India."),
        ("Web access", "Available as a public website with no login needed for searching."),
    ]:
        check(doc, label, text)
    subheading(doc, "Limitations of the Existing System:", indent=0.0)
    for text in [
        "Search works one district at a time; there is no nationwide or 'near me' view.",
        "No blood-group compatibility help: a family searching O- is not told which donor groups are safe.",
        "No indication of how fresh a record is, so months-old stock looks the same as today's.",
        "No history, trends or shortage view for health officials.",
        "No dedicated search for rare groups such as the Bombay (Oh) phenotype.",
        "No data export for reports and planning.",
    ]:
        bullet(doc, text, indent=0.6)

    heading(doc, "Proposed System")
    para(
        doc,
        "Raktsetu is a web platform built on a data lakehouse. It collects eRaktKosh stock for every "
        "district, cleans and validates it, keeps a daily history, and presents it through an "
        "emergency-first, mobile-friendly website. It adds clinical compatibility rules, freshness "
        "and trust indicators, distance ranking and analytics that the existing portal does not offer.",
        indent=0.3,
        after=10,
    )
    subheading(doc, "Key Features", indent=0.0)
    features = [
        (
            "Smart Emergency Mode",
            [
                "Enter the patient's blood group and component; the system lists compatible donor groups (NBTC rules, including plasma and Bombay rules).",
                "Ranks blood centres by freshness, trust, units and distance, with one-tap Call and Directions.",
            ],
        ),
        (
            "Blood Finder & Live Map",
            [
                "Filter by state, district, component, blood group and name; sort by freshness, units, trust or distance.",
                "'Near me' uses the device location to show the closest centres on a map.",
            ],
        ),
        (
            "Freshness, Trust Score & 7-day Trend",
            [
                "Every card shows when the stock was last updated and a 0–100 Trust Score.",
                "A trend chart shows real daily snapshots; days without data appear as gaps, never invented values.",
            ],
        ),
        (
            "Rare Blood & District Shortage Radar",
            [
                "Nationwide list of Bombay and rare-group stock.",
                "District-level shortage severity for each blood group, exportable as CSV.",
            ],
        ),
        (
            "Donor Eligibility Checker",
            [
                "Checks age, weight, haemoglobin, recent tattoo/surgery/illness and donation gap in the browser only.",
                "Shows the next eligible donation date; no personal data leaves the device.",
            ],
        ),
        (
            "Data Health Dashboard & PWA",
            [
                "Shows last sync time, fresh vs stale centres, dbt test results and recent pipeline runs.",
                "Installable, responsive and usable offline with the last cached data.",
            ],
        ),
    ]
    for title, items in features:
        check(doc, title)
        for item in items:
            bullet(doc, item)

    comparison(doc)
    heading(doc, "Scope of the Proposed System")
    para(
        doc,
        "Raktsetu covers the full path from public eRaktKosh data to an emergency decision: data "
        "collection, quality control, history, analytics and a public website. It is an information "
        "service; it does not reserve or transfer blood, and users are always told to call the blood "
        "centre before travelling.",
        indent=0.3,
        after=8,
    )
    subheading(
        doc,
        "Public User Side Scope (patients' families, doctors, donors, officials)",
        indent=0.0,
        size=12,
    )
    for i, (t, d) in enumerate(
        [
            (
                "Emergency Search",
                "Find compatible blood for a patient and call the best-ranked centres.",
            ),
            (
                "Blood Finder",
                "Search and filter all centres, view stock per blood group, switch to the map.",
            ),
            ("Near Me", "Use the device location to rank centres by distance."),
            ("Trends & Trust", "View a centre's recent stock history and its Trust Score."),
            (
                "Rare Blood & Shortages",
                "Find rare-group stock and districts that are running short.",
            ),
            ("Reports", "Download availability and shortage data as CSV files."),
            ("Donor Help", "Check donation eligibility and the next eligible date."),
        ],
        1,
    ):
        numbered(doc, i, t, d)
    subheading(doc, "Operator (Admin) Side Scope", indent=0.0, size=12)
    for i, (t, d) in enumerate(
        [
            ("Scrape", "Run the nationwide crawl that writes an immutable Parquet file."),
            (
                "Load & Transform",
                "Load the file into Bronze and build Silver/Gold tables with dbt.",
            ),
            ("Quality Gate", "Run 37 data tests and the facility-count anomaly check."),
            ("Publish", "Atomically release the new data to the website."),
            ("Monitor", "Track sync age, stale centres and pipeline runs on the Data Health page."),
        ],
        1,
    ):
        numbered(doc, i, t, d)

    heading(doc, "Aim and Objectives")
    p = para(doc, "", indent=0.3, after=10)
    _font(p.add_run("Aim: "), 12, bold=True)
    _font(
        p.add_run(
            "To build a reliable, data-driven web platform that converts India's public eRaktKosh "
            "blood-stock data into fast, trustworthy and clinically aware guidance for people searching "
            "for blood in an emergency, and into shortage intelligence for health officials."
        ),
        12,
    )
    para(doc, "Objectives", bold=True, indent=0.3, after=4)
    for text in [
        "Collect nationwide blood stock automatically and politely from every eRaktKosh district.",
        "Store raw data immutably and build cleaned, tested Bronze, Silver and Gold layers.",
        "Keep a daily snapshot history so trends can be shown honestly over time.",
        "Apply NBTC compatibility rules for red cells, plasma and the Bombay phenotype.",
        "Rank blood centres using freshness, trust, available units and distance.",
        "Show district-level shortages and rare-blood availability, with CSV export.",
        "Block publishing when data quality checks fail, so users never see broken data.",
        "Deliver a responsive, accessible, installable website that also works offline.",
        "Protect privacy: no login, no personal data stored, location never saved.",
    ]:
        check(doc, text)

    heading(doc, "Feasibility Study")
    heading(doc, "Operational Feasibility", page_break=False)
    para(
        doc,
        "Operational feasibility examines how well the solution fits the way people already search for "
        "blood and how easily it can be operated.",
        indent=0.3,
        after=8,
    )
    for label, text in [
        (
            "User Roles",
            "Public users search without any account. The operator runs one command (make pipeline) to refresh data.",
        ),
        (
            "User Experience (UX)",
            "Emergency search needs only three choices and one click; results show Call and Directions buttons directly.",
        ),
        (
            "Training",
            "No training is needed for public users. The operator needs a short guide covering the pipeline commands.",
        ),
        (
            "Integration with Existing Process",
            "Raktsetu does not replace eRaktKosh; it reads its public data and adds compatibility, ranking and analytics on top.",
        ),
        (
            "Support & Maintenance",
            "Automated tests, the anomaly gate and the Data Health page make problems visible early.",
        ),
        (
            "Risk Factors",
            "If eRaktKosh changes or fails, the circuit breaker and quality gate stop the run and the last good data stays online.",
        ),
    ]:
        check(doc, label, text)

    heading(doc, "Technical Feasibility")
    for label, text in [
        (
            "Technology Stack",
            "Python, FastAPI, DuckDB, dbt and vanilla web technologies are mature, open-source and well documented.",
        ),
        (
            "Platform Compatibility",
            "Runs in any modern browser on Android, iOS, Windows, macOS and Linux; installable as a PWA.",
        ),
        (
            "Performance",
            "DuckDB answers analytical queries over about 66,000 current stock rows in milliseconds on a laptop.",
        ),
        (
            "Reliability",
            "Read-only serving database and atomic publishing keep the website consistent during updates.",
        ),
        (
            "Scalability & Maintainability",
            "Modular code (scraper, pipeline, API, web) and dbt models can be extended independently.",
        ),
        (
            "Development Tools",
            "uv, Ruff, pytest, Playwright, Docker and Git provide a reproducible, tested build.",
        ),
    ]:
        check(doc, label, text)

    heading(doc, "Economical Feasibility")
    para(
        doc,
        "Economic feasibility determines whether the project is cost-effective and what value it returns.",
        indent=0.3,
        after=8,
    )
    for label, text in [
        (
            "Development Cost",
            "All software used is free and open source; development was done by the two-member team.",
        ),
        (
            "Initial Investment",
            "No licences or paid APIs; the data source and map geocoding are free public services.",
        ),
        (
            "Operational Cost",
            "A single low-cost virtual server can host the API and run the daily pipeline.",
        ),
        (
            "Cost-Benefit Analysis",
            "Saves families time during emergencies and gives officials shortage reports without extra data entry.",
        ),
        (
            "Social Return on Investment",
            "Faster access to compatible blood can directly reduce delays in treatment.",
        ),
        (
            "Scalability Cost",
            "Costs grow slowly: DuckDB handles millions of rows on one machine before any upgrade is needed.",
        ),
    ]:
        check(doc, label, text)


# ============================================================================
# 4. System analysis
# ============================================================================
def system_analysis(doc):
    divider(doc, "SYSTEM ANALYSIS")
    heading(doc, "Requirement Specification (along with System Modules)")
    modules = [
        (
            "Data Ingestion Module",
            [
                "Async scrape of all districts listed in the eRaktKosh metadata catalogue.",
                "Retries with back-off, circuit breaker, minimum 6-hour run interval.",
                "Writes an immutable Parquet file plus a manifest with SHA-256 checksum.",
            ],
        ),
        (
            "Lakehouse Transformation Module",
            [
                "Idempotent load of Parquet into bronze.raw_inventory.",
                "dbt models: Silver staging views and Gold facts, dimensions and marts.",
                "Incremental daily snapshots; groups missing from the latest scrape count as 0 units.",
            ],
        ),
        (
            "Quality Gate & Publishing Module",
            [
                "37 dbt data tests and an 80% facility-count anomaly gate.",
                "Run log in gold.pipeline_runs; atomic copy to serving.duckdb.",
            ],
        ),
        (
            "Emergency Module",
            [
                "Compatible donor groups for red cells, plasma and Bombay phenotype.",
                "Weighted ranking (freshness 35%, trust 30%, units 25%, distance 10%); radius search near the user.",
            ],
        ),
        (
            "Blood Finder & Map Module",
            ["Filters, sorting, pagination, URL-shareable state, grid and map views, 'Near me'."],
        ),
        (
            "Analytics Module",
            [
                "National distribution chart, 7-day trend, Trust Score, rare blood radar, district shortages."
            ],
        ),
        (
            "Donor Assistance Module",
            [
                "Client-side eligibility rules (NBTC) with next eligible date; no data sent to the server."
            ],
        ),
        (
            "Data Health & Export Module",
            [
                "Sync age, fresh/stale counts, dbt results, recent runs; CSV export of availability and shortages."
            ],
        ),
        (
            "REST API Module",
            [
                "Documented endpoints under /api/v1 with validation, OpenAPI docs and health check (/healthz)."
            ],
        ),
    ]
    for i, (name, items) in enumerate(modules, 1):
        p = para(doc, "", align=WD_ALIGN_PARAGRAPH.LEFT, before=4, after=2)
        _font(p.add_run(f"{i}. {name}:"), 12, bold=True)
        for item in items:
            check(doc, item)
    subheading(doc, "Non-Functional Requirements:", indent=0.0)
    for label, text in [
        ("Performance", "Search responses within about one second for nationwide queries."),
        (
            "Privacy",
            "No login, no personal data stored; donor check and location stay in the browser.",
        ),
        ("Reliability", "Website keeps serving the last good data if a pipeline run fails."),
        (
            "Usability & Accessibility",
            "Responsive from 320 px, 40 px touch targets, keyboard-closable dialogs.",
        ),
    ]:
        check(doc, label, text)

    heading(doc, "Hardware and Software Requirements")
    para(doc, "Server / Developer Side", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=2)
    table(
        doc,
        ["Item", "Minimum Requirement"],
        [
            ("Processor", "Dual-core 2 GHz or better"),
            ("RAM", "4 GB (8 GB recommended)"),
            ("Storage", "10 GB free disk space"),
            ("Operating System", "Linux, Windows 10/11 or macOS"),
            ("Internet", "Required for scraping eRaktKosh and map tiles"),
            ("Software", "Python 3.12, uv, dbt-duckdb, FastAPI, Uvicorn, Git; Docker (optional)"),
        ],
        (2.0, 4.3),
        caption="Server and developer requirements",
    )
    para(doc, "Client (User) Side", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=2)
    table(
        doc,
        ["Item", "Minimum Requirement"],
        [
            ("Device", "Any smartphone, tablet, laptop or desktop"),
            ("Screen", "320 px wide or larger"),
            ("Browser", "Latest Chrome, Firefox, Edge or Safari"),
            ("Internet", "Needed for live data; the last data stays available offline"),
            ("Location", "Optional, only for 'Near me' (needs HTTPS or localhost)"),
        ],
        (2.0, 4.3),
        caption="Client requirements",
    )

    diagram(
        doc,
        "Data Flow Diagram (Level 0)",
        "The context diagram shows Raktsetu as one process and the outside entities that send data to it "
        "or receive data from it.",
        "dfd_level0.png",
    )
    diagram(
        doc,
        "Data Flow Diagram (Level 1)",
        "Level 1 breaks the system into six processes and shows the four data stores between them.",
        "dfd_level1.png",
    )
    diagram(
        doc,
        "Use Case Diagram: Public Users",
        "Four kinds of public users use the website without logging in. «include» means the step always "
        "happens; «extend» means it is optional.",
        "use_case_diagram.png",
    )
    heading(doc, "Actors of the System")
    table(
        doc,
        ["Actor", "Type", "What they do"],
        [
            (
                "Patient's Family",
                "Primary",
                "Searches for compatible blood in an emergency and calls the centre",
            ),
            (
                "Doctor / Hospital Staff",
                "Primary",
                "Checks compatible groups, rare blood and stock trends",
            ),
            ("Health Official", "Primary", "Monitors district shortages and exports reports"),
            ("Blood Donor", "Primary", "Checks donation eligibility and where blood is needed"),
            ("Operator (Admin)", "Primary", "Runs the data pipeline and monitors data health"),
            ("eRaktKosh Portal", "Secondary (system)", "Provides blood stock data"),
            ("OpenStreetMap", "Secondary (system)", "Provides district coordinates and map data"),
            ("Device GPS", "Secondary (system)", "Provides the user's location for 'Near me'"),
        ],
        (1.9, 1.5, 2.9),
        caption="Actors of the system",
    )
    diagram(
        doc,
        "Use Case Diagram: Operator & External Systems",
        "The operator runs the pipeline, which always includes scraping, loading, transforming, testing and "
        "publishing. eRaktKosh and OpenStreetMap are external systems.",
        "use_case_operator.png",
    )
    diagram(
        doc,
        "Class Diagram",
        "Main classes of the system with their attributes and operations, and how they depend on each other.",
        "class_diagram.png",
    )
    diagram(
        doc,
        "Activity Diagram: Public User",
        "How a user finds blood: emergency search or Blood Finder, then Call or Directions.",
        "activity_diagram.png",
    )
    diagram(
        doc,
        "Activity Diagram: Data Pipeline",
        "One pipeline run from scraping to publishing, including the checks that can stop it.",
        "activity_diagram_pipeline.png",
    )
    diagram(
        doc,
        "Sequence Diagram",
        "Messages exchanged between the user, website, API, medical rules and database during an "
        "emergency search.",
        "sequence_diagram.png",
    )
    diagram(
        doc,
        "System Architecture",
        "End-to-end design: from the eRaktKosh portal, through the lakehouse layers and quality gate, to the "
        "website.",
        "architecture_diagram.png",
    )


# ============================================================================
# 5. System design
# ============================================================================
DESCRIPTIONS = {
    "run_id": "Unique identifier of the pipeline/scrape run",
    "scraped_at": "Date-time when the record was scraped",
    "hospital_code": "Unique eRaktKosh code of the blood centre",
    "hospital_name": "Name of the blood centre",
    "state_name": "State or Union Territory",
    "district_name": "District of the blood centre",
    "hospital_type": "Ownership type (Govt, Private, Charitable)",
    "facility_type": "Facility category (e.g. Blood Bank)",
    "hospital_contact": "Phone, fax and e-mail as published",
    "hospital_address": "Postal address of the blood centre",
    "entry_date_raw": "Last stock update time as text from the source",
    "components_raw": "Stock per component and blood group (JSON)",
    "raw_payload": "Complete original record (JSON) for audit",
    "ingested_at": "Date-time when loaded into Bronze",
    "source_updated_at": "Last stock update time reported by the centre",
    "lakehouse_ingested_at": "Time of the latest scrape containing the centre",
    "component_name": "Blood component (e.g. Packed Red Blood Cells)",
    "blood_group": "Normalised blood group (e.g. O-, AB+, Bombay)",
    "units": "Units available",
    "snapshot_at": "Scrape time of the snapshot",
    "snapshot_date": "Calendar date of the daily snapshot",
    "last_synced_at": "Time of the centre's latest scrape",
    "total_reporting_banks": "Number of centres reporting this group",
    "stocked_banks": "Centres with at least one unit",
    "banks_with_zero": "Centres with zero units",
    "total_units": "Total units in the district",
    "min_facility_units": "Lowest stock at a single centre",
    "max_facility_units": "Highest stock at a single centre",
    "shortage_flag": "True when total units are 5 or fewer",
    "severity_score": "Shortage urgency from 0 to 100",
    "shortage_status": "Critical depletion / High / Moderate / Adequate",
    "rarity_tier": "Rarity class of the blood group",
    "days_since_update": "Days since the centre last updated stock",
    "distinct_components_count": "Number of components the centre reports",
    "total_units_in_stock": "Total units across all components",
    "freshness_points": "Trust points for recent updates (max 50)",
    "contact_points": "Trust points for a valid phone (max 25)",
    "completeness_points": "Trust points for component coverage (max 25)",
    "reliability_score": "Trust Score from 0 to 100",
    "reliability_label": "High (75+), Medium (40–74) or Low",
    "reliability_tier": "Same as label, kept for API compatibility",
    "current_units": "Units in the latest snapshot",
    "prev_snapshot_units": "Units in the previous snapshot",
    "delta_prev": "Change since the previous snapshot",
    "delta_7d": "Change over the last 7 snapshots",
    "risk_flag": "Stable, Falling or Depleted",
    "anomaly_flag": "NORMAL or SUDDEN_DROP_TO_ZERO",
    "started_at": "Run start time",
    "finished_at": "Run finish time",
    "status": "RUNNING, SUCCESS, SUCCESS_NO_SCRAPE or FAILED",
    "facilities_count": "Blood centres in the published data",
    "tests_passed": "Whether all quality checks passed",
    "error_message": "Failure reason, if any",
    "created_at": "Row creation time",
}

DATA_DICTIONARY = [
    ("1. Bronze Layer (Raw Data)", [("1.1 raw_inventory", "bronze.raw_inventory")]),
    (
        "2. Gold Layer: Dimension & Facts",
        [
            ("2.1 dim_blood_banks", "main_gold.dim_blood_banks"),
            ("2.2 fct_stock_snapshot", "main_gold.fct_stock_snapshot"),
            ("2.3 fct_stock_current", "main_gold.fct_stock_current"),
        ],
    ),
    (
        "3. Gold Layer: Analytical Marts",
        [
            ("3.1 mart_district_status", "main_gold.mart_district_status"),
            ("3.2 mart_rare_blood_availability", "main_gold.mart_rare_blood_availability"),
            ("3.3 mart_reliability_score", "main_gold.mart_reliability_score"),
            ("3.4 mart_stock_change", "main_gold.mart_stock_change"),
        ],
    ),
    ("4. Metadata", [("4.1 pipeline_runs", "gold.pipeline_runs")]),
]


def _columns(table_name):
    import duckdb

    db = os.path.join(REPO_ROOT, "serving.duckdb")
    con = duckdb.connect(db, read_only=True)
    try:
        return [(r[0], r[1]) for r in con.execute(f"DESCRIBE {table_name}").fetchall()]
    finally:
        con.close()


def system_design(doc):
    divider(doc, "SYSTEM DESIGN")
    diagram(
        doc,
        "ER Diagram",
        "Entities stored in the lakehouse and how they are related. PK = primary key, FK = foreign key; "
        "underlined fields together identify a row.",
        "er_diagram.png",
    )
    first = True
    for group, tables in DATA_DICTIONARY:
        if first:
            heading(doc, "Data Dictionary")
            first = False
        else:
            new_page(doc)
        para(doc, group, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=4, indent=0.2)
        for title, name in tables:
            rows = [
                (c, t.replace("HUGEINT", "BIGINT"), DESCRIPTIONS.get(c, ""))
                for c, t in _columns(name)
            ]
            table(
                doc,
                ["Field Name", "Data Type", "Description"],
                rows,
                (2.2, 1.3, 3.0),
                font_size=10,
                caption=f"{title.split(' ', 1)[1]} ({name})",
            )

    divider(doc, "Screen Layouts")
    for title, text, shot in [
        (
            "Home Page & National Statistics",
            "The first screen shows live national figures: total blood centres, "
            "centres updated in the last 24 hours and centres flagged for review.",
            "01_home.png",
        ),
        (
            "Smart Emergency Mode",
            "The user picks the patient's blood group, component and place. The system shows "
            "safe donor groups, warnings for Rh-negative and Bombay groups, and the best centres with Call and "
            "Directions buttons.",
            "02_emergency.png",
        ),
        (
            "Blood Finder",
            "Filters for state, district, component, blood group and name, with sorting by freshness, "
            "units, trust or distance.",
            "03_finder.png",
        ),
        (
            "Blood Finder: Result Cards",
            "Each card shows the centre, its location, how recently it updated, its "
            "Trust Score and the units for all eight blood groups.",
            "03b_cards.png",
        ),
        (
            "Live Blood Map",
            "Search results on a map. Marker colour shows the stock level; a pop-up offers Call, "
            "Directions and the 7-day trend.",
            "04_map.png",
        ),
        (
            "7-day Inventory Trend",
            "Real daily stock history per blood group, with current units, change and status.",
            "05_trend_modal.png",
        ),
        (
            "Trust Score Explanation",
            "How the 0–100 Trust Score is calculated: freshness (50), valid contact (25) "
            "and component coverage (25).",
            "06_trust_modal.png",
        ),
        (
            "National Network Distribution",
            "Bar chart of active blood centres in the top states.",
            "07_network_chart.png",
        ),
        (
            "Rare Blood Phenotype Radar",
            "Centres across India holding Bombay and other rare blood, with contact numbers.",
            "08_rare_blood.png",
        ),
        (
            "District Shortage Radar",
            "Districts running short of the selected blood group, with CSV export for "
            "health officials.",
            "09_shortages.png",
        ),
        (
            "Donor Eligibility & Assistance",
            "Checks in the browser whether a person can donate today, using NBTC rules.",
            "10_donor.png",
        ),
        (
            "Data Health & Pipeline Runs",
            "Last successful sync, fresh and stale centres, test results and recent pipeline runs.",
            "11_data_health.png",
        ),
    ]:
        screen(doc, title, text, shot)
    for title, text, pair in [
        (
            "Mobile Responsive Design",
            "The same website on a phone: home screen and the menu.",
            ("12_mobile_home.png", "13_mobile_menu.png"),
        ),
        (
            "Mobile Responsive Design: Search Results",
            "Blood centre cards and emergency results on a phone, with "
            "large Call and Directions buttons.",
            ("14_mobile_cards.png", "15_mobile_emergency.png"),
        ),
    ]:
        heading(doc, title)
        check(doc, text)
        images_side_by_side(
            doc, [os.path.join(SCREENSHOTS_DIR, f) for f in pair], width_each=2.9, caption=title
        )


# ============================================================================
# 6. System testing
# ============================================================================
TEST_CASES = [
    (
        "Home: Load Dashboard",
        "Verify live statistics load",
        "Open the website",
        "Total, fresh and flagged centre counts shown",
    ),
    (
        "Finder: State → District",
        "Verify dependent dropdowns",
        "Select state Gujarat",
        "District list is filled and enabled",
    ),
    (
        "Finder: District Search",
        "Verify filters apply",
        "Gujarat, Anand, Search",
        "Only Anand centres listed; URL updated",
    ),
    (
        "Finder: Blood Group Filter",
        "Verify group filter",
        "Click B-",
        "Centres with B- stock listed",
    ),
    (
        "Finder: No Results",
        "Verify empty state",
        "Search 'zzzz-no-such-bank'",
        "'No matching blood centres' message",
    ),
    (
        "Finder: Reset",
        "Verify filters reset",
        "Click Reset",
        "All filters cleared; all centres listed",
    ),
    (
        "Finder: Sort by Units",
        "Verify sorting",
        "Select 'Available Units'",
        "Highest stock shown first",
    ),
    (
        "Finder: Pagination",
        "Verify page navigation",
        "Click Next, Last, change page size",
        "Correct page ranges shown",
    ),
    (
        "Finder: Near Me",
        "Verify location ranking",
        "Allow location (Bengaluru)",
        "Bangalore Urban centres about 4 km away first",
    ),
    (
        "Card: Call Link",
        "Verify phone link",
        "Tap Call on a card",
        "Opens dialer with the first phone number only",
    ),
    (
        "Card: 7-day Trend",
        "Verify history modal",
        "Click 7-day trend, change component",
        "Chart and summary table update",
    ),
    (
        "Card: Trust Score",
        "Verify explanation",
        "Click Trust badge",
        "Score method shown; closes with Esc",
    ),
    (
        "Map: View & Pop-up",
        "Verify map view",
        "Switch to map, click marker",
        "Markers and pop-up with actions shown",
    ),
    (
        "Emergency: O- Patient",
        "Verify compatibility",
        "Group O-, Search",
        "Only O- donors; Rh-negative warning shown",
    ),
    (
        "Emergency: Bombay",
        "Verify rare rule",
        "Group Bombay (Oh+)",
        "Only Bombay donors; critical warning shown",
    ),
    (
        "Emergency: AB+ Plasma",
        "Verify plasma rule",
        "AB+, Fresh Frozen Plasma",
        "Only AB donor groups listed",
    ),
    (
        "Emergency: Invalid Group",
        "Verify validation",
        "API call with group 'XYZ'",
        "Error: Unknown blood group",
    ),
    (
        "Shortages: Filters",
        "Verify shortage radar",
        "State Gujarat, group AB-",
        "Only Gujarat AB- districts listed",
    ),
    (
        "Export: Availability CSV",
        "Verify CSV format",
        "Click Export CSV",
        "Valid CSV, header on first row, all rows",
    ),
    (
        "Export: Shortages CSV",
        "Verify full export",
        "Export O- nationwide",
        "All 501 districts exported",
    ),
    ("Donor: Eligible", "Verify eligibility", "Age 25, 70 kg, 100 days", "'Likely Eligible' shown"),
    ("Donor: Under Age", "Verify deferral", "Age 17", "'Temporarily Deferred' with reason"),
    (
        "Donor: Future Date",
        "Verify date validation",
        "Last donation in future",
        "Clear message to correct the date",
    ),
    ("Data Health", "Verify monitoring", "Open Data Health", "Sync age, runs and dbt tests shown"),
    (
        "Mobile Layout",
        "Verify responsiveness",
        "Open at 320–390 px width",
        "No horizontal scroll; menu button works",
    ),
    (
        "Offline Mode",
        "Verify PWA cache",
        "Load once, go offline, reload",
        "App opens with offline notice",
    ),
    (
        "Pipeline: Quality Gate",
        "Verify anomaly gate",
        "Facility count below 80%",
        "Publish blocked; old data stays online",
    ),
    (
        "Pipeline: Repeat Load",
        "Verify idempotent load",
        "Load the same Parquet twice",
        "Second load skipped; no duplicates",
    ),
]


def system_testing(doc):
    divider(doc, "SYSTEM TESTING")
    heading(doc, "Testing Strategies")
    for label, text in [
        (
            "Black Box Testing",
            "Every user flow was tested from the outside by using the real website in a browser "
            "(Playwright), on desktop and mobile screen sizes.",
        ),
        (
            "White Box (Unit) Testing",
            "pytest tests check the code inside: blood-group rules, Trust Score, "
            "geocoding, CSV format and the scraper safety checks.",
        ),
        (
            "Integration Testing",
            "Every API endpoint is called with FastAPI's TestClient against a test database.",
        ),
        (
            "Data Quality Testing",
            "dbt tests check that keys are unique, required values exist and values are "
            "valid in every layer. The anomaly gate checks the facility count.",
        ),
        (
            "Responsive Testing",
            "The layout was checked at 320, 360, 390, 768, 1024, 1280 and 1920 px widths.",
        ),
        (
            "Offline (PWA) Testing",
            "The website was reloaded with the network turned off to check the cache.",
        ),
    ]:
        check(doc, label, text)

    heading(doc, "Test Results Summary")
    table(
        doc,
        ["Test Type", "Tool", "Number of Checks", "Result"],
        [
            ("Unit & integration tests", "pytest", "98", "All passed"),
            ("Data quality tests", "dbt", "37", "All passed"),
            ("End-to-end user flows", "Playwright", "146", "All passed"),
            ("Responsive layout", "Playwright", "7 screen widths", "No horizontal scrolling"),
            ("Black box test cases", "Manual + automated", str(len(TEST_CASES)), "All passed"),
        ],
        (2.1, 1.4, 1.4, 1.4),
        center_cols=(2,),
        caption="Summary of test results",
    )

    heading(doc, "Test Cases")
    para(doc, "Black box test cases: each row was checked on the running website.", after=6)
    rows = [(str(i), s, p, inp, out, "PASS") for i, (s, p, inp, out) in enumerate(TEST_CASES, 1)]
    table(
        doc,
        ["Sr. No.", "Scenario", "Purpose", "Input", "Expected Output", "Status"],
        rows,
        (0.45, 1.35, 1.2, 1.35, 1.6, 0.55),
        header_fill="D9D9D9",
        font_size=10,
        center_cols=(0, 5),
        caption="Black box test cases",
    )


# ============================================================================
# 7. Conclusion
# ============================================================================
def conclusion(doc):
    divider(doc, "CONCLUSION")
    heading(doc, "Conclusion")
    for text in [
        "Raktsetu shows that open government data can be turned into a practical emergency tool. The system "
        "collects blood stock for 4,716 blood centres across 35 States/UTs, cleans and tests it, keeps a daily "
        "history and serves it through a fast website that works on any phone.",
        "Compared with searching the eRaktKosh portal one district at a time, a user can now find compatible "
        "blood near them in a few seconds, see how fresh each record is and call the centre with one tap. "
        "Health officials get district-level shortage reports that can be downloaded as CSV files.",
        "The project also gave us hands-on experience with modern data engineering: an async scraper, a medallion "
        "lakehouse with DuckDB and dbt, automated data-quality gates, a REST API and a responsive Progressive "
        "Web App, all backed by automated tests.",
    ]:
        para(doc, text, after=10)
    heading(doc, "Limitations")
    for text in [
        "Data is a periodic snapshot of eRaktKosh, not real-time; users must call before travelling.",
        "Accuracy depends on blood centres updating their stock on eRaktKosh.",
        "Locations are district centres, not exact hospital addresses.",
        "History covers only the days the pipeline has run, so long-term trends need more time.",
        "'Near me' works only on HTTPS or localhost because browsers protect location access.",
        "The interface is available in English only.",
        "This is an academic project, not an official government service.",
    ]:
        check(doc, text)


def future_enhancement(doc):
    divider(doc, "FUTURE ENHANCEMENT")
    heading(doc, "Future Enhancement")
    groups = [
        (
            "Data & Pipeline",
            [
                (
                    "Scheduled Scraping",
                    "Run the pipeline automatically every 6–12 hours to build a longer daily history.",
                ),
                (
                    "Hospital-level Geocoding",
                    "Locate each blood centre from its address instead of the district centre.",
                ),
                (
                    "Time-travel Tables",
                    "Apache Iceberg or Delta Lake to query stock as it was on any past date.",
                ),
                (
                    "Advanced Data Quality",
                    "Great Expectations checks for unusual values before publishing.",
                ),
            ],
        ),
        (
            "Analytics & Machine Learning",
            [
                (
                    "Stock Forecasting",
                    "Predict when a blood group will run out at a centre (e.g. 'O- in about 3 days').",
                ),
                (
                    "Shortage Early Warning",
                    "Flag districts trending towards zero, not only those already at zero.",
                ),
                ("Transfer Suggestions", "Recommend moving surplus units between nearby centres."),
                (
                    "Seasonal Analysis",
                    "Study stock changes during dengue season, festivals and holidays.",
                ),
            ],
        ),
        (
            "User Features",
            [
                (
                    "Stock Alerts",
                    "E-mail or push notification when a needed group becomes available nearby.",
                ),
                (
                    "Blood Centre Portal",
                    "Let centres confirm or correct their own stock to raise trust.",
                ),
                ("Donation Camp Calendar", "Show camps in districts with shortages."),
                ("Multilingual Interface", "Hindi and Gujarati with medically reviewed wording."),
                ("Voice Search & Accessibility", "Voice input and a full WCAG 2.2 AA audit."),
            ],
        ),
        (
            "Deployment & Operations",
            [
                ("HTTPS Cloud Hosting", "Public deployment so 'Near me' works on every phone."),
                ("CI/CD", "GitHub Actions running all tests and dbt checks on every change."),
                (
                    "Monitoring & Alerts",
                    "Prometheus/Grafana dashboards and failure alerts for the pipeline.",
                ),
                ("Public API Keys", "Rate-limited API access for hospitals and research groups."),
            ],
        ),
    ]
    for i, (group, items) in enumerate(groups):
        if i:
            heading(doc, f"Future Enhancement: {group}")
        else:
            subheading(doc, group + ":", indent=0.0)
        for n, (title, text) in enumerate(items, 1):
            numbered(doc, n, title, text)
    heading(doc, "Long-Term Expansion Opportunities")
    for label, text in [
        (
            "National Integration",
            "Partner with state blood transfusion councils for direct data feeds.",
        ),
        (
            "Hospital Information Systems",
            "Offer the API to hospital software for automatic stock lookups.",
        ),
        ("Inter-state Coordination", "Shortage maps to support movement of blood between states."),
        ("Research Data", "Anonymised historical datasets for public-health research."),
    ]:
        check(doc, label, text)


def references(doc):
    divider(doc, "REFERENCES")
    heading(doc, "References")
    for url in [
        "https://eraktkosh.mohfw.gov.in/",
        "https://nbtc.mohfw.gov.in/",
        "https://fastapi.tiangolo.com/",
        "https://duckdb.org/docs/",
        "https://docs.getdbt.com/",
        "https://leafletjs.com/reference.html",
        "https://www.chartjs.org/docs/latest/",
        "https://nominatim.org/release-docs/latest/api/Search/",
        "https://carto.com/basemaps",
        "https://developer.mozilla.org/en-US/docs/Web/Progressive_web_apps",
        "https://playwright.dev/python/",
        "https://docs.pytest.org/",
        "https://github.com/radhey803/eraktkosh-blood-lakehouse",
    ]:
        p = para(doc, "", after=4, indent=0.6, align=WD_ALIGN_PARAGRAPH.LEFT)
        _font(p.add_run("✓ "), 12)
        _font(p.add_run(url), 12, underline=True, color=RGBColor(0x05, 0x63, 0xC1))


def thank_you(doc):
    new_page(doc)
    para(
        doc,
        "THANK YOU",
        size=36,
        bold=True,
        underline=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=290,
    )


# ============================================================================
# Build (two passes for the Table of Contents)
# ============================================================================
def build(path, pages, captions=None):
    """Builds the report; returns the figure and table captions it created."""
    FIGS.clear()
    TABS.clear()
    figs, tabs = captions or ([], [])
    doc = docx.Document()
    setup_document(doc)
    cover(doc)
    declaration(doc)
    acknowledgement(doc)
    abstract(doc)
    table_of_contents(doc, pages)
    list_page(doc, "List of Figures", "Figure", figs, pages)
    list_page(doc, "List of Tables", "Table", tabs, pages)
    FIGS.clear()
    TABS.clear()
    project_profile(doc)
    project_timeline(doc)
    introduction_to_tools(doc)
    system_study(doc)
    system_analysis(doc)
    system_design(doc)
    system_testing(doc)
    conclusion(doc)
    future_enhancement(doc)
    references(doc)
    appendix(doc)
    thank_you(doc)
    doc.save(path)
    return list(FIGS), list(TABS)


def to_pdf(docx_path, out_dir):
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", out_dir, docx_path],
        check=True,
        capture_output=True,
    )
    return os.path.join(out_dir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")


def find_pages(pdf_path, figs, tabs):
    """Maps TOC subjects and 'Figure N:' / 'Table N:' captions to their PDF page numbers."""
    text = subprocess.run(
        ["pdftotext", "-layout", pdf_path, "-"], check=True, capture_output=True, text=True
    ).stdout
    page_lines = [
        [ln.strip() for ln in page.splitlines() if ln.strip()] for page in text.split("\f")
    ]
    toc_page = next(i for i, lines in enumerate(page_lines) if "Table of Contents" in lines)
    pages, start = {}, toc_page + 1
    body_start = None
    for _, _, subject, anchor, _ in TOC:
        if not anchor:
            continue
        for i in range(start, len(page_lines)):
            lines = page_lines[i]
            words = " ".join(lines).split()
            on_divider = len(words) < 12 and anchor in " ".join(words)
            if on_divider or any(
                ln.startswith(anchor) and len(ln) <= len(anchor) + 4 for ln in lines
            ):
                pages[subject] = i + 1
                start = i
                if anchor == "PROJECT PROFILE":
                    body_start = i
                break
    for kind, items in (("Figure", figs), ("Table", tabs)):
        for n in range(1, len(items) + 1):
            key = f"{kind} {n}:"
            for i in range(body_start or toc_page, len(page_lines)):
                if any(ln.startswith(key) for ln in page_lines[i]):
                    pages[key] = i + 1
                    break
    return pages


def main():
    if not shutil.which("soffice") or not shutil.which("pdftotext"):
        raise SystemExit(
            "LibreOffice (soffice) and poppler (pdftotext) are needed to number the TOC."
        )
    os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        draft = os.path.join(tmp, "draft.docx")
        print("Pass 1: collecting figure and table captions...")
        captions = build(draft, {})
        print("Pass 2: locating page numbers...")
        captions = build(draft, {}, captions)
        pages = find_pages(to_pdf(draft, tmp), *captions)
        missing = [s for _, _, s, a, _ in TOC if a and s not in pages]
        if missing:
            print("Warning: page not found for", missing)
    print("Pass 3: building final report...")
    build(OUTPUT_FILE, pages, captions)
    pdf = to_pdf(OUTPUT_FILE, os.path.dirname(OUTPUT_FILE))
    print(f"Saved {OUTPUT_FILE}\nSaved {pdf}")
    print(
        f"{len(captions[0])} figures, {len(captions[1])} tables; "
        f"{sum(k.startswith(('Figure', 'Table')) for k in pages)} caption pages located"
    )


if __name__ == "__main__":
    main()
